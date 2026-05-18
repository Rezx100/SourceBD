"""US DoL ILAB List of Goods Produced by Child or Forced Labor (TVPRA list).

Source listing page (Tier 5 regulatory):
  https://www.dol.gov/agencies/ilab/reports/child-labor/list-of-goods

The Bureau of International Labor Affairs publishes the dataset as an XLSX
attached to the listing page. Filename pattern observed:
  /sites/dolgov/files/ILAB/child_labor_reports/tda{YYYY}/{YYYY}-TVPRA-list-of-goods.xlsx

This scraper:
  1. Fetches the listing page HTML.
  2. Picks the most-recent .xlsx whose href contains 'tvpra' or
     'list-of-goods'.
  3. Downloads that workbook.
  4. Iterates rows in the first sheet. Standard column layout:
       Country | Good | Child Labor | Forced Labor | Forced Child Labor
     (cells are Y/N or year ranges). We collapse the three exploitation flags
     into a single status string.
  5. Yields one SanctionEntry per (country, good) row.

These rows describe goods/sectors, not specific entities, so they will not
match any supplier under our 2-significant-token gate (false-positive defence).
They are still ingested to power "audited against ILAB" coverage badges and
the country/merchandise-level risk lookup tables we surface on supplier
profile pages.
"""
from __future__ import annotations

import io
import re
from typing import AsyncIterator
from urllib.parse import urljoin

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from slugify import slugify

from etl.core.http import HttpClient
from etl.core.sanctions import BaseSanctionScraper, SanctionEntry

LIST_PAGE = "https://www.dol.gov/agencies/ilab/reports/child-labor/list-of-goods"

_XLSX_HINT_RE = re.compile(r"(tvpra|list[-_]of[-_]goods)", re.IGNORECASE)
_YEAR_RE = re.compile(r"(20\d{2})")


def _pick_xlsx_url(html: str, base_url: str) -> str | None:
    soup = BeautifulSoup(html, "lxml")
    candidates: list[tuple[int, str]] = []
    for a in soup.find_all("a", href=True):
        href = a["href"]
        if not href.lower().endswith(".xlsx"):
            continue
        if not _XLSX_HINT_RE.search(href):
            continue
        url = urljoin(base_url, href)
        m = _YEAR_RE.search(url)
        year = int(m.group(1)) if m else 0
        candidates.append((year, url))
    if not candidates:
        return None
    candidates.sort(reverse=True)
    return candidates[0][1]


def _norm_header(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def _is_present(cell) -> bool:
    """ILAB marks exploitation type with 'Y', 'X', a year range, or similar
    non-empty text. Treat any non-empty, non-'N' cell as positive."""
    if cell is None:
        return False
    v = str(cell).strip()
    if not v:
        return False
    if v.upper() in ("N", "NO", "NA", "N/A", "-", "—"):
        return False
    return True


class IlabTvpraScraper(BaseSanctionScraper):
    code = "ilab_tvpra"
    source_code = "ILAB"

    async def fetch(self) -> AsyncIterator[SanctionEntry]:
        async with HttpClient(rps=0.5) as http:
            page = await http.get(LIST_PAGE)
            xlsx_url = _pick_xlsx_url(page.text, LIST_PAGE)
            if not xlsx_url:
                self.log.error("ilab.no_xlsx_link", listing_page=LIST_PAGE)
                return
            self.log.info("ilab.xlsx", url=xlsx_url)
            xlsx_resp = await http.get(xlsx_url)
            xlsx_bytes = xlsx_resp.content

        wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]

        rows = ws.iter_rows(values_only=True)
        # Find the header row: first row containing both 'country' and 'good'.
        # ILAB column headers seen in the wild: 'Country/Area', 'Good',
        # 'Child Labor', 'Forced Labor' (and sometimes 'Forced Child Labor').
        col_country = col_good = -1
        col_child = col_forced = col_forced_child = -1
        for raw_row in rows:
            if raw_row is None:
                continue
            cells = [_norm_header(c) if c is not None else "" for c in raw_row]
            if any(c.startswith("country") for c in cells) and any("good" in c for c in cells):
                for i, c in enumerate(cells):
                    if c.startswith("country") and col_country < 0:
                        col_country = i
                    elif "good" in c and col_good < 0:
                        col_good = i
                    elif "forced child" in c:
                        col_forced_child = i
                    elif "forced labor" in c:
                        col_forced = i
                    elif "child labor" in c:
                        col_child = i
                break

        if col_country < 0 or col_good < 0:
            self.log.error("ilab.header_not_found", url=xlsx_url)
            return

        for raw_row in rows:
            if raw_row is None:
                continue
            country = (raw_row[col_country] if col_country < len(raw_row) else None)
            good = (raw_row[col_good] if col_good < len(raw_row) else None)
            if not country or not good:
                continue
            country = str(country).strip()
            good = str(good).strip()
            if not country or not good:
                continue

            flags: list[str] = []
            if col_child >= 0 and col_child < len(raw_row) and _is_present(raw_row[col_child]):
                flags.append("Child Labor")
            if col_forced >= 0 and col_forced < len(raw_row) and _is_present(raw_row[col_forced]):
                flags.append("Forced Labor")
            if (
                col_forced_child >= 0
                and col_forced_child < len(raw_row)
                and _is_present(raw_row[col_forced_child])
            ):
                flags.append("Forced Child Labor")
            status = ", ".join(flags) if flags else "Listed"

            ref = f"ilab-{slugify(country)}-{slugify(good)}"

            yield SanctionEntry(
                list_code="ilab_tvpra",
                source_code="ILAB",
                entry_ref=ref,
                entity_name=f"{country} — {good}",
                aliases=[],
                country=country,
                merchandise=good,
                listed_date=None,
                status=status,
                status_notes=None,
                source_url=xlsx_url,
                raw={"country": country, "good": good, "flags": flags},
            )
